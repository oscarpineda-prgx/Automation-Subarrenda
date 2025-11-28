"""
Utilidad para aplicar formato a los Excel generados (detalle y resumen).

- Convierte columnas de fecha a formato corto mm/dd/yyyy.
- Aplica formato monetario a columnas de importes.
- Colorea encabezados con fondo azul oscuro y texto blanco en negritas.
"""

import datetime as _dt
import os
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

# Palabras clave para detectar columnas monetarias
_MONEY_KEYS = [
    "importe",
    "subtotal",
    "total",
    "mtto",
    "renta",
    "diferencia",
    "dif",
    "auditoria",
]


def _looks_money(header: str) -> bool:
    h = header.lower()
    return any(k in h for k in _MONEY_KEYS)


def _looks_date(header: str) -> bool:
    h = header.lower()
    return "fecha" in h or "f_contrato" in h


def _to_datetime(val):
    """Convierte strings o date a datetime para que Excel lo trate como fecha."""
    if val is None or val == "":
        return None
    if isinstance(val, _dt.datetime):
        return val
    if isinstance(val, _dt.date):
        return _dt.datetime.combine(val, _dt.time.min)
    if isinstance(val, (int, float)):
        # Dejar números (podrían ser fechas ya numéricas en Excel)
        return val
    if isinstance(val, str):
        txt = val.strip()
        # Intentos comunes
        for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%d/%m/%Y", "%m/%d/%Y"):
            try:
                return _dt.datetime.strptime(txt, fmt)
            except ValueError:
                continue
        try:
            return _dt.datetime.fromisoformat(txt)
        except ValueError:
            return val
    return val


def _to_number(val):
    """Convierte strings con signos a número flotante."""
    if val is None or val == "":
        return None
    if isinstance(val, (int, float)):
        return val
    if isinstance(val, str):
        txt = val.replace("$", "").replace(",", "").replace("(", "-").replace(")", "").strip()
        try:
            return float(txt)
        except ValueError:
            return val
    return val


def _style_headers(headers: Iterable):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(color="FFFFFF", bold=True)
    align = Alignment(horizontal="center", vertical="center")
    for cell in headers:
        cell.fill = fill
        cell.font = font
        cell.alignment = align


def format_workbook(path: str):
    wb = load_workbook(path)
    for ws in wb.worksheets:
        # Oculta líneas de cuadrícula para una apariencia limpia
        ws.sheet_view.showGridLines = False

        if ws.max_row < 2:
            _style_headers(ws[1])
            continue

        headers = [cell.value or "" for cell in ws[1]]
        # Detecta columnas por palabra clave en el encabezado
        money_cols = {idx for idx, h in enumerate(headers, start=1) if _looks_money(str(h))}
        date_cols = {idx for idx, h in enumerate(headers, start=1) if _looks_date(str(h))}

        _style_headers(ws[1])

        for row in ws.iter_rows(min_row=2):
            for idx, cell in enumerate(row, start=1):
                if idx in date_cols:
                    new_val = _to_datetime(cell.value)
                    cell.value = new_val
                    if isinstance(new_val, (_dt.datetime, _dt.date)):
                        cell.number_format = "mm/dd/yyyy"
                elif idx in money_cols:
                    num = _to_number(cell.value)
                    cell.value = num
                    if isinstance(num, (int, float)):
                        # Formato Accounting en Excel
                        cell.number_format = '_("$"* #,##0.00_);_("$"* (#,##0.00);_("$"* "-"??_);_(@_)'
    wb.save(path)


def main():
    files = [
        os.path.join("data", "output", "detalle_subarrendatarios.xlsx"),
        os.path.join("data", "output", "resumen_subarrendatarios.xlsx"),
    ]
    for f in files:
        if os.path.isfile(f):
            format_workbook(f)
            print(f"Formato aplicado a {f}")
        else:
            print(f"Archivo no encontrado: {f} (omitido)")


if __name__ == "__main__":
    main()
