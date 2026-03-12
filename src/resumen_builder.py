import sys
import pandas as pd
import os

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.drawing.image import Image

from src.format_excel import format_workbook

# Permitir imports desde src
sys.path.append("src")
from loader import load_all
from detalle_builder import generar_detalle_todos
from detalle_exporter import (
    _buscar_mes_fda,
    _buscar_mes_fda_por_orden,
    _coerce_mes,
    _mes_a_texto,
)


def generar_resumen(
    detalle: pd.DataFrame,
    df_clientes: pd.DataFrame = None,
    df_contratos: pd.DataFrame = None,
) -> pd.DataFrame:
    """
    Genera un resumen por contrato.
    - Si existe columna `orden`, agrupa por (orden, rfc, subarrendatario).
    - Si no existe, mantiene compatibilidad agrupando por (rfc, subarrendatario).
    - Suma diferencia_base_vs_aud.
    - Agrega sumas de mtto/subtotales/totales solo sobre filas con cobro (importe_renta_c > 0).
    - Incluye cuota de mantenimiento (%).
    - Incluye el primer importe_renta_c > 0 y el primer importe_renta_auditoria > 0 (o 0 si no existen).
    - Marca si hay cobro (importe_renta_c > 0 en alguna fila) o NO HAY COBRO.
    - No colapsa contratos con mismo RFC/cliente cuando existe `orden`.
    """
    if detalle.empty:
        return pd.DataFrame(
            columns=[
                "orden",
                "rfc",
                "subarrendatario",
                "area",
                "f_contrato_ini",
                "f_contrato_fin",
                "renta_mensual",
                "duracion_meses",
                "cuota_mantenimiento",
                "sum_importe_mtto_c",
                "sum_importe_mtto_a",
                "sum_subtotal_c",
                "sum_subtotal_a",
                "sum_total_c",
                "sum_total_a",
                "sum_dif_base_vs_aud",
                "primer_importe_renta_c",
                "primer_importe_renta_a",
                "estatus_cobro",
                "acta_entrega",
                "count_meses_sin_cobro",
            ]
        )

    group_cols = ["rfc", "subarrendatario"]
    if "orden" in detalle.columns:
        group_cols = ["orden", "rfc", "subarrendatario"]

    def _primer_importe(df_base: pd.DataFrame, columna: str) -> pd.Series:
        """
        Devuelve serie con el primer valor > 0 por contrato (group_cols) para la columna indicada.
        Si no hay valores > 0, devuelve 0.
        """
        if columna not in df_base.columns:
            return pd.Series(dtype=float)

        ordenado = (
            df_base.sort_values("fecha")
            if "fecha" in df_base.columns
            else df_base.copy()
        )
        serie = (
            ordenado[ordenado[columna] > 0]
            .groupby(group_cols)[columna]
            .first()
            .round(2)
        )
        return serie

    base_primeros = detalle.copy()
    primeros_cliente = _primer_importe(base_primeros, "importe_renta_c")
    primeros_auditoria = _primer_importe(base_primeros, "importe_renta_auditoria")

    def _sumas_condicion_cobro(df_base: pd.DataFrame) -> pd.DataFrame:
        """
        Suma importes solo en filas con cobro (importe_renta_c > 0).
        Devuelve dataframe con columnas de sumas y cuota de mantenimiento.
        """
        key_cols = group_cols
        cols = key_cols + [
            "importe_mtto_c_cobro_sum",
            "importe_mtto_auditoria_cobro_sum",
            "subtotal_c_cobro_sum",
            "subtotal_a_cobro_sum",
            "total_c_cobro_sum",
            "total_a_cobro_sum",
            "cuota_mantenimiento_pct",
        ]
        if df_base.empty:
            return pd.DataFrame(columns=cols)

        claves = df_base[key_cols].drop_duplicates()

        con_cobro = df_base[df_base["importe_renta_c"].fillna(0) > 0].copy()
        sumas = (
            con_cobro.assign(
                importe_mtto_c=lambda d: pd.to_numeric(d.get("importe_mtto_c", 0), errors="coerce").fillna(0),
                importe_mtto_auditoria=lambda d: pd.to_numeric(d.get("importe_mtto_auditoria", 0), errors="coerce").fillna(0),
                subtotal_c=lambda d: pd.to_numeric(d.get("subtotal_c", 0), errors="coerce").fillna(0),
                subtotal_a=lambda d: pd.to_numeric(d.get("subtotal_a", 0), errors="coerce").fillna(0),
                total_c=lambda d: pd.to_numeric(d.get("total_c", 0), errors="coerce").fillna(0),
                total_a=lambda d: pd.to_numeric(d.get("total_a", 0), errors="coerce").fillna(0),
            )
            .groupby(group_cols)
            .agg(
                importe_mtto_c_cobro_sum=("importe_mtto_c", "sum"),
                importe_mtto_auditoria_cobro_sum=("importe_mtto_auditoria", "sum"),
                subtotal_c_cobro_sum=("subtotal_c", "sum"),
                subtotal_a_cobro_sum=("subtotal_a", "sum"),
                total_c_cobro_sum=("total_c", "sum"),
                total_a_cobro_sum=("total_a", "sum"),
            )
            .reset_index()
        )

        # Ensamblar con todas las claves rfc+subarrendatario
        sumas = claves.merge(sumas, on=key_cols, how="left")

        # Cuota por contrato (no depende de que haya cobro)
        if "cuota_mantenimiento" in df_base.columns:
            cuota_df = (
                df_base.assign(cuota_mantenimiento=lambda d: pd.to_numeric(d["cuota_mantenimiento"], errors="coerce"))
                .groupby(key_cols, as_index=False)["cuota_mantenimiento"]
                .first()
                .rename(columns={"cuota_mantenimiento": "cuota_mantenimiento_pct"})
            )
            sumas = sumas.merge(cuota_df, on=key_cols, how="left")
        else:
            sumas["cuota_mantenimiento_pct"] = None

        for col in cols:
            if col not in sumas.columns:
                sumas[col] = None
        return sumas[cols]

    sumas_cobro = _sumas_condicion_cobro(detalle)

    def _acta_entrega(df_base: pd.DataFrame) -> pd.Series:
        """
        Marca si el incremento del cliente ocurre en el mismo mes de aniversario que el de auditoria.
        - Innecesaria: ambos importes incrementan en el primer mes de aniversario.
        - Necesaria: auditoria incrementa en aniversario pero cliente no en ese mes.
        - Desconocida: faltan datos para evaluar.
        """
        req_cols = {"fecha", "importe_renta_c", "importe_renta_auditoria"}
        if not req_cols.issubset(df_base.columns):
            return pd.Series(dtype=object)

        def evaluar(grp: pd.DataFrame) -> str:
            g = grp.copy()
            g["fecha"] = pd.to_datetime(g["fecha"], errors="coerce")
            g = g.dropna(subset=["fecha"]).sort_values("fecha")
            if g.empty:
                return "Desconocida"
            inicio = g.iloc[0]["fecha"]
            if pd.isna(inicio):
                return "Desconocida"
            aniversario = g[(g["fecha"] > inicio) & (g["fecha"].dt.month == inicio.month)]
            if aniversario.empty:
                return "Desconocida"
            ann = aniversario.iloc[0]
            ann_idx = g.index.get_loc(ann.name)
            if isinstance(ann_idx, slice) or ann_idx is None:
                return "Desconocida"
            prev_pos = ann_idx - 1
            if prev_pos < 0:
                return "Desconocida"
            prev = g.iloc[prev_pos]
            cliente_prev = pd.to_numeric(prev["importe_renta_c"], errors="coerce")
            cliente_ann = pd.to_numeric(ann["importe_renta_c"], errors="coerce")
            aud_prev = pd.to_numeric(prev["importe_renta_auditoria"], errors="coerce")
            aud_ann = pd.to_numeric(ann["importe_renta_auditoria"], errors="coerce")
            if any(pd.isna([cliente_prev, cliente_ann, aud_prev, aud_ann])):
                return "Desconocida"
            inc_cliente = cliente_ann > cliente_prev
            inc_auditoria = aud_ann > aud_prev
            if inc_auditoria and inc_cliente:
                return "Innecesaria"
            if inc_auditoria and not inc_cliente:
                return "Necesaria"
            return "Desconocida"

        res = df_base.groupby(group_cols).apply(evaluar)
        res.name = "acta_entrega"
        return res

    acta = _acta_entrega(detalle)

    def _count_meses_sin_cobro(df_base: pd.DataFrame) -> pd.Series:
        """
        Cuenta meses sin cobro dentro del rango entre el primer y ultimo importe_renta_c > 0.
        Si nunca hay cobro, devuelve 0.
        """
        if "importe_renta_c" not in df_base.columns:
            return pd.Series(dtype=int)

        def contar(grp: pd.DataFrame) -> int:
            ordenado = grp.sort_values("fecha") if "fecha" in grp.columns else grp.copy()
            vals = pd.to_numeric(ordenado["importe_renta_c"], errors="coerce").fillna(0)
            pos = vals[vals > 0]
            if pos.empty:
                return 0
            start = pos.index[0]
            end = pos.index[-1]
            rango = vals.loc[start:end]
            return int((rango <= 0).sum())

        res = df_base.groupby(group_cols).apply(contar)
        res.name = "count_meses_sin_cobro"
        return res

    count_meses = _count_meses_sin_cobro(detalle)
    # Info de contrato derivada del detalle (sin agregar columnas nuevas al detalle)
    agregados = {}
    if "area" in detalle.columns:
        agregados["area"] = ("area", "first")
    if "fecha" in detalle.columns:
        agregados["fecha_contrato_inicio"] = ("fecha", "min")
        agregados["fecha_contrato_fin"] = ("fecha", "max")
        agregados["duracion_meses_contrato"] = ("fecha", "count")
    if "importe_renta_auditoria" in detalle.columns:
        agregados["renta_mensual_contrato"] = ("importe_renta_auditoria", "first")

    if agregados:
        contrato_info = (
            detalle.groupby(group_cols)
            .agg(**agregados)
            .reset_index()
        )
    else:
        contrato_info = pd.DataFrame(columns=group_cols)

    # Base completa de claves rfc + subarrendatario (no se eliminan duplicados de RFC con nombres distintos)
    claves = detalle[group_cols].drop_duplicates()

    def _sumar_diferencia_en_rango(
        df_base: pd.DataFrame, claves_df: pd.DataFrame
    ) -> pd.DataFrame:
        """
        Suma diferencia_base_vs_aud entre la primera y la ultima fecha con cobro (>0) por contrato.
        Si faltan fechas validas, cae al comportamiento previo (solo suma filas con cobro).
        """
        req_cols = {"importe_renta_c", "diferencia_base_vs_aud"}
        if df_base.empty or not req_cols.issubset(df_base.columns):
            res = claves_df.copy()
            res["diferencia_base_vs_aud_sum"] = 0
            return res

        if "fecha" not in df_base.columns:
            df_dif = df_base[df_base["importe_renta_c"].fillna(0) > 0].copy()
            diferencia = (
                df_dif.fillna({"diferencia_base_vs_aud": 0})
                .groupby(group_cols, as_index=False)["diferencia_base_vs_aud"]
                .sum()
                .rename(
                    columns={"diferencia_base_vs_aud": "diferencia_base_vs_aud_sum"}
                )
            )
            return (
                claves_df.merge(diferencia, on=group_cols, how="left")
                .fillna({"diferencia_base_vs_aud_sum": 0})
            )

        def sumar(grp: pd.DataFrame) -> float:
            g = grp.copy()
            g["fecha"] = pd.to_datetime(g["fecha"], errors="coerce")
            importes = pd.to_numeric(g["importe_renta_c"], errors="coerce").fillna(0)
            cobros = importes > 0
            if not cobros.any():
                return 0.0

            cobros_validos = g.loc[cobros & g["fecha"].notna()]
            if cobros_validos.empty:
                dif_cobro = pd.to_numeric(
                    g.loc[cobros, "diferencia_base_vs_aud"], errors="coerce"
                ).fillna(0)
                return float(dif_cobro.sum())

            fecha_inicio = cobros_validos["fecha"].min()
            fecha_fin = cobros_validos["fecha"].max()
            en_rango = (g["fecha"] >= fecha_inicio) & (g["fecha"] <= fecha_fin)
            diferencias = pd.to_numeric(
                g["diferencia_base_vs_aud"], errors="coerce"
            ).fillna(0)
            return float(diferencias[en_rango.fillna(False)].sum())

        sumas = (
            df_base.groupby(group_cols)
            .apply(sumar)
            .reset_index(name="diferencia_base_vs_aud_sum")
        )
        return (
            claves_df.merge(sumas, on=group_cols, how="left")
            .fillna({"diferencia_base_vs_aud_sum": 0})
        )

    # Agrupa diferencia dentro del rango entre el primer y ultimo cobro (>0) y mantiene todas las claves.
    diferencia_sum = _sumar_diferencia_en_rango(detalle, claves)

    def _incremento(df_base: pd.DataFrame) -> pd.Series:
        """Toma el primer valor no nulo de incremento por contrato."""
        if "incremento" not in df_base.columns:
            return pd.Series(dtype=object)

        def pick(grp: pd.DataFrame):
            vals = grp["incremento"].dropna()
            if vals.empty:
                return None
            return vals.iloc[0]

        res = df_base.groupby(group_cols).apply(pick)
        res.name = "incremento"
        return res

    incremento = _incremento(detalle)

    def _calcular_incrementos(
        df_base: pd.DataFrame,
        df_cli: pd.DataFrame,
        df_con: pd.DataFrame,
    ) -> pd.DataFrame:
        """
        Obtiene mes incremento FDA/AUD y diferencia en meses por contrato (mismos criterios que detalle_exporter).
        """
        cols = group_cols + [
            "mes_incremento_fda",
            "mes_incremento_aud",
            "diferencia_meses",
        ]
        if df_base.empty:
            return pd.DataFrame(columns=cols)

        registros = []
        for _, grp in df_base.groupby(group_cols):
            rfc_val = str(grp["rfc"].iloc[0]).strip() if "rfc" in grp.columns else ""
            sub_val = (
                str(grp["subarrendatario"].iloc[0]).strip()
                if "subarrendatario" in grp.columns
                else ""
            )
            if "orden" in grp.columns:
                ord_val = pd.to_numeric(grp["orden"].iloc[0], errors="coerce")
                mes_fda_num, mes_fda_txt = _buscar_mes_fda_por_orden(
                    df_cli, ord_val, rfc_val, sub_val
                )
            else:
                mes_fda_num, mes_fda_txt = _buscar_mes_fda(
                    df_cli, rfc_val, sub_val
                )
            mes_aud_num = None
            mes_aud_txt = None
            if "mes_incremento" in grp.columns:
                for val in grp["mes_incremento"]:
                    mes_val = _coerce_mes(val)
                    if mes_val:
                        mes_aud_num = mes_val
                        mes_aud_txt = _mes_a_texto(mes_val)
                        break

            diferencia = (
                abs(int(mes_fda_num) - int(mes_aud_num))
                if mes_fda_num is not None and mes_aud_num is not None
                else None
            )
            fila = {c: grp[c].iloc[0] for c in group_cols}
            fila.update(
                {
                    "mes_incremento_fda": mes_fda_txt,
                    "mes_incremento_aud": mes_aud_txt,
                    "diferencia_meses": diferencia,
                }
            )
            registros.append(fila)

        return pd.DataFrame(registros, columns=cols)

    incrementos = _calcular_incrementos(detalle, df_clientes, df_contratos)
    # Bandera de cobro: si existe al menos un importe_renta_c > 0
    estatus = (
        detalle.fillna({"importe_renta_c": 0})
        .assign(tiene_cobro=lambda d: d["importe_renta_c"] > 0)
        .groupby(group_cols)["tiene_cobro"]
        .any()
        .replace({True: "SI HAY COBRO", False: "NO HAY COBRO"})
    )

    resumen = diferencia_sum
    resumen = resumen.merge(
        primeros_cliente.rename("primer_importe_renta_c").reset_index(),
        on=group_cols,
        how="left",
    )
    resumen = resumen.merge(
        primeros_auditoria.rename("primer_importe_renta_auditoria").reset_index(),
        on=group_cols,
        how="left",
    )
    resumen["primer_importe_renta_c"] = resumen[
        "primer_importe_renta_c"
    ].fillna(0)
    resumen["primer_importe_renta_auditoria"] = resumen[
        "primer_importe_renta_auditoria"
    ].fillna(0)
    resumen = resumen.merge(
        estatus.rename("estatus_cobro").reset_index(), on=group_cols, how="left"
    )
    resumen["estatus_cobro"] = resumen["estatus_cobro"].fillna("NO HAY COBRO")

    resumen = resumen.merge(
        incremento.reset_index(), on=group_cols, how="left"
    )
    resumen["incremento"] = resumen["incremento"].fillna("INPC")

    resumen = resumen.merge(
        acta.reset_index(), on=group_cols, how="left"
    )
    resumen["acta_entrega"] = resumen["acta_entrega"].fillna("Desconocida")

    resumen = resumen.merge(
        count_meses.rename("count_meses_sin_cobro").reset_index(),
        on=group_cols,
        how="left",
    )
    resumen["count_meses_sin_cobro"] = resumen["count_meses_sin_cobro"].fillna(0).astype(int)

    if not incrementos.empty:
        resumen = resumen.merge(incrementos, on=group_cols, how="left")
    else:
        resumen["mes_incremento_fda"] = None
        resumen["mes_incremento_aud"] = None
        resumen["diferencia_meses"] = None

    resumen = resumen.merge(sumas_cobro, on=group_cols, how="left")
    for c in [
        "importe_mtto_c_cobro_sum",
        "importe_mtto_auditoria_cobro_sum",
        "subtotal_c_cobro_sum",
        "subtotal_a_cobro_sum",
        "total_c_cobro_sum",
        "total_a_cobro_sum",
    ]:
        if c in resumen.columns:
            resumen[c] = resumen[c].fillna(0).round(2)
    if "cuota_mantenimiento_pct" in resumen.columns:
        resumen["cuota_mantenimiento_pct"] = resumen["cuota_mantenimiento_pct"].round(4)

    if not contrato_info.empty:
        resumen = resumen.merge(contrato_info, on=group_cols, how="left")

    # Renombrar columnas a los alias solicitados por el usuario
    column_renames = {
        "fecha_contrato_inicio": "f_contrato_ini",
        "fecha_contrato_fin": "f_contrato_fin",
        "renta_mensual_contrato": "renta_mensual",
        "duracion_meses_contrato": "duracion_meses",
        "cuota_mantenimiento_pct": "cuota_mantenimiento",
        "importe_mtto_c_cobro_sum": "sum_importe_mtto_c",
        "subtotal_c_cobro_sum": "sum_subtotal_c",
        "total_c_cobro_sum": "sum_total_c",
        "primer_importe_renta_auditoria": "primer_importe_renta_a",
        "importe_mtto_auditoria_cobro_sum": "sum_importe_mtto_a",
        "subtotal_a_cobro_sum": "sum_subtotal_a",
        "total_a_cobro_sum": "sum_total_a",
        "diferencia_base_vs_aud_sum": "sum_dif_base_vs_aud",
    }
    resumen = resumen.rename(columns=column_renames)

    # Reordenar columnas segun solicitud
    orden = [
        "orden",
        "rfc",
        "subarrendatario",
        "area",
        "f_contrato_ini",
        "f_contrato_fin",
        "renta_mensual",
        "duracion_meses",
        "cuota_mantenimiento",
        "primer_importe_renta_c",
        "sum_importe_mtto_c",
        "sum_subtotal_c",
        "sum_total_c",
        "primer_importe_renta_a",
        "sum_importe_mtto_a",
        "sum_subtotal_a",
        "sum_total_a",
        "sum_dif_base_vs_aud",
        "incremento",
        "estatus_cobro",
        "acta_entrega",
        "count_meses_sin_cobro",
        "mes_incremento_fda",
        "mes_incremento_aud",
        "diferencia_meses",
    ]
    presentes = [c for c in orden if c in resumen.columns]
    resto = [c for c in resumen.columns if c not in presentes]
    resumen = resumen[presentes + resto]

    return resumen


def generar_resumen_desde_fuente() -> pd.DataFrame:
    """
    Carga clientes/contratos/INPC, genera el detalle completo y devuelve el resumen,
    sin leer/escribir archivos intermedios.
    """
    clientes, contratos, inpc = load_all()
    detalle = generar_detalle_todos(contratos, clientes, inpc)
    return generar_resumen(detalle, clientes, contratos)


def aplicar_presentacion_resumen(
    ruta: str,
    logo_path: str = "data/input/Picture1.png",
) -> None:
    """Inserta logo y titulos en la parte superior de la hoja."""
    wb = load_workbook(ruta)
    ws = wb.active

    # Reservar espacio arriba del encabezado original
    espacio = 7
    ws.insert_rows(1, espacio)

    # Logo a la izquierda
    try:
        if os.path.isfile(logo_path):
            img = Image(logo_path)
            img.width = 180
            img.height = 80
            ws.add_image(img, "B2")
        else:
            print(f"Logo no encontrado en {logo_path}")
    except Exception as exc:
        print(f"No se pudo insertar logo en {ruta}: {exc}")

    titulo = ws.cell(row=4, column=6, value="RESUMEN SUBARRENDATARIOS - DIFERENCIA DE COBRO")
    titulo.font = Font(bold=True, size=14)
    titulo.alignment = Alignment(horizontal="center")
    ws.merge_cells(start_row=4, start_column=6, end_row=4, end_column=13)

    wb.save(ruta)


if __name__ == "__main__":
    resumen = generar_resumen_desde_fuente()
    output_path = "data/output/resumen_subarrendatarios.xlsx"
    resumen.to_excel(output_path, index=False)
    aplicar_presentacion_resumen(output_path)
    print(f"Resumen guardado en {output_path}")
